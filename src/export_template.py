# export_template.py
# 按模板导出 - 自动识别 + 手动修正版
# 修改：导出支持“另存为”（选择路径和文件名）
# 修复：导出设置窗口左右分栏布局

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import pandas as pd
from pathlib import Path
from datetime import datetime
import os
import re
from copy import copy
from openpyxl import load_workbook

from field_matcher import FIELD_DISPLAY_NAMES


def detect_template_structure(file_path):
    """
    自动识别模板结构
    返回: (header_row, data_start_row, total_row, audit_row)
    所有行号均为 0-based（从0开始）
    """
    try:
        df_raw = pd.read_excel(file_path, header=None, engine='openpyxl')
    except Exception:
        return None, None, None, None
    
    if df_raw.empty:
        return None, None, None, None
    
    header_keywords = ['固定资产名称', '取得时间', '使用年限', '固定资产原值', '残值率']
    header_row = None
    
    for idx in range(min(30, len(df_raw))):
        row_text = ' '.join([str(cell) for cell in df_raw.iloc[idx] if pd.notna(cell)])
        match_count = sum(1 for kw in header_keywords if kw in row_text)
        if match_count >= 2:
            header_row = idx
            break
    
    if header_row is None:
        best_count = 0
        for idx in range(min(30, len(df_raw))):
            non_empty = df_raw.iloc[idx].notna().sum()
            if non_empty > best_count:
                best_count = non_empty
                header_row = idx
        if header_row is None or best_count < 3:
            return None, None, None, None
    
    data_start_row = header_row + 1
    
    total_row = None
    for idx in range(header_row + 1, len(df_raw)):
        row_text = ' '.join([str(cell) for cell in df_raw.iloc[idx] if pd.notna(cell)])
        if '合计' in row_text:
            total_row = idx
            break
    
    audit_row = None
    for idx in range(header_row + 1, len(df_raw)):
        row_text = ' '.join([str(cell) for cell in df_raw.iloc[idx] if pd.notna(cell)])
        if '审计说明' in row_text:
            audit_row = idx
            break
    
    if audit_row is None and total_row is not None:
        audit_row = total_row + 2
    
    return header_row, data_start_row, total_row, audit_row


def prepare_merged_data(df, dep_result):
    """合并原始数据和计算结果"""
    if df is None or df.empty:
        return pd.DataFrame()
    
    data = df.copy()
    
    if dep_result is not None and not dep_result.empty:
        dep_copy = dep_result.copy()
        
        join_col = None
        if '资产编号' in data.columns and '资产编号' in dep_copy.columns:
            join_col = '资产编号'
        elif 'asset_id' in data.columns and 'asset_id' in dep_copy.columns:
            join_col = 'asset_id'
        elif '资产编号' in data.columns and 'asset_id' in dep_copy.columns:
            dep_copy.rename(columns={'asset_id': '资产编号'}, inplace=True)
            join_col = '资产编号'
        elif 'asset_id' in data.columns and '资产编号' in dep_copy.columns:
            dep_copy.rename(columns={'资产编号': 'asset_id'}, inplace=True)
            join_col = 'asset_id'
        else:
            common = set(data.columns).intersection(set(dep_copy.columns))
            if common:
                join_col = list(common)[0]
            else:
                return data
        
        data[join_col] = data[join_col].astype(str)
        dep_copy[join_col] = dep_copy[join_col].astype(str)
        data = pd.merge(data, dep_copy, on=join_col, how='left', suffixes=('', '_dep'))
        
        to_drop = [c for c in data.columns if c.endswith('_dep') and c[:-4] in data.columns]
        data.drop(columns=to_drop, inplace=True, errors='ignore')
    
    return data


def copy_row_style(ws, source_row, target_row, max_col):
    """从 source_row 复制样式到 target_row"""
    for col in range(1, max_col + 1):
        source_cell = ws.cell(row=source_row, column=col)
        target_cell = ws.cell(row=target_row, column=col)
        if source_cell.has_style:
            try:
                target_cell.font = copy(source_cell.font)
                target_cell.border = copy(source_cell.border)
                target_cell.fill = copy(source_cell.fill)
                target_cell.number_format = source_cell.number_format
                target_cell.alignment = copy(source_cell.alignment)
            except Exception:
                pass


class ExportSettingsWindow:
    def __init__(self, parent, template_path, df, dep_result, base_date):
        self.parent = parent
        self.template_path = template_path
        self.df = df
        self.dep_result = dep_result
        self.base_date = base_date
        self.result = None
        self.initialized = False
        
        self.window = tk.Toplevel(parent)
        self.window.title("导出设置：按模板导出")
        self.window.geometry("850x560")
        self.window.minsize(800, 500)
        self.window.transient(parent)
        self.window.grab_set()
        
        header_row_0, data_start_row_0, total_row_0, audit_row_0 = detect_template_structure(template_path)
        
        if header_row_0 is None:
            ttk.Label(self.window, text="无法识别模板表头，请确认模板格式", foreground='red').pack(pady=50)
            ttk.Button(self.window, text="关闭", command=self.window.destroy).pack(pady=10)
            return
        
        self.default_header = header_row_0 + 1
        self.default_data_start = data_start_row_0 + 1
        self.default_total = total_row_0 + 1 if total_row_0 is not None else None
        self.default_audit = audit_row_0 + 1 if audit_row_0 is not None else None
        
        try:
            df_fields = pd.read_excel(template_path, header=header_row_0, engine='openpyxl')
            self.template_fields = [col for col in df_fields.columns if pd.notna(col) and str(col).strip() != '']
        except Exception:
            self.template_fields = []
        
        if not self.template_fields:
            ttk.Label(self.window, text="无法从模板中提取字段", foreground='red').pack(pady=50)
            ttk.Button(self.window, text="关闭", command=self.window.destroy).pack(pady=10)
            return
        
        self.merged_data = prepare_merged_data(df, dep_result)
        if self.merged_data.empty:
            ttk.Label(self.window, text="数据合并失败，请检查资产编号", foreground='red').pack(pady=50)
            ttk.Button(self.window, text="关闭", command=self.window.destroy).pack(pady=10)
            return
        
        self.source_fields = []
        for col in self.merged_data.columns:
            display = FIELD_DISPLAY_NAMES.get(col, col)
            if display not in self.source_fields and display != '':
                self.source_fields.append(display)
        
        if not self.source_fields:
            self.source_fields = list(self.merged_data.columns)
        
        self.initialized = True
        self._build_ui()
        
        self.window.update_idletasks()
        x = (self.window.winfo_screenwidth() - 850) // 2
        y = (self.window.winfo_screenheight() - 560) // 2
        self.window.geometry(f"+{x}+{y}")
    
    def _build_ui(self):
        info_frame = ttk.Frame(self.window)
        info_frame.pack(fill=tk.X, padx=10, pady=8)
        
        ttk.Label(info_frame, text=f"📄 模板文件: {Path(self.template_path).name}", font=('Microsoft YaHei', 10)).pack(anchor=tk.W)
        
        struct_frame = ttk.LabelFrame(self.window, text="模板结构确认（程序已自动识别，可手动修改）", padding=8)
        struct_frame.pack(fill=tk.X, padx=10, pady=5)
        
        row_frame1 = ttk.Frame(struct_frame)
        row_frame1.pack(fill=tk.X, pady=2)
        ttk.Label(row_frame1, text="表头行:").pack(side=tk.LEFT, padx=5)
        self.header_entry = ttk.Entry(row_frame1, width=6)
        self.header_entry.insert(0, str(self.default_header))
        self.header_entry.pack(side=tk.LEFT, padx=5)
        
        ttk.Label(row_frame1, text="数据起始行:").pack(side=tk.LEFT, padx=15)
        self.data_start_entry = ttk.Entry(row_frame1, width=6)
        self.data_start_entry.insert(0, str(self.default_data_start))
        self.data_start_entry.pack(side=tk.LEFT, padx=5)
        
        row_frame2 = ttk.Frame(struct_frame)
        row_frame2.pack(fill=tk.X, pady=2)
        ttk.Label(row_frame2, text="合计行:").pack(side=tk.LEFT, padx=5)
        self.total_entry = ttk.Entry(row_frame2, width=6)
        self.total_entry.insert(0, str(self.default_total) if self.default_total else '')
        self.total_entry.pack(side=tk.LEFT, padx=5)
        
        ttk.Label(row_frame2, text="审计说明行:").pack(side=tk.LEFT, padx=15)
        self.audit_entry = ttk.Entry(row_frame2, width=6)
        self.audit_entry.insert(0, str(self.default_audit) if self.default_audit else '')
        self.audit_entry.pack(side=tk.LEFT, padx=5)
        
        ttk.Label(row_frame2, text="（填写数字，空表示未识别）", foreground='gray').pack(side=tk.LEFT, padx=10)
        
        paned = ttk.PanedWindow(self.window, orient=tk.HORIZONTAL)
        paned.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        left_frame = ttk.Frame(paned, width=400)
        paned.add(left_frame, weight=1)
        
        ttk.Label(left_frame, text="字段映射（将模板列与数据源字段关联）", font=('Microsoft YaHei', 9, 'bold')).pack(anchor=tk.W, pady=(0, 5))
        
        table_container = ttk.Frame(left_frame)
        table_container.pack(fill=tk.BOTH, expand=True)
        
        canvas = tk.Canvas(table_container, highlightthickness=0, height=320)
        scrollbar = ttk.Scrollbar(table_container, orient=tk.VERTICAL, command=canvas.yview)
        canvas.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        self.mapping_frame = ttk.Frame(canvas)
        canvas.create_window((0, 0), window=self.mapping_frame, anchor='nw')
        
        def _on_configure(event):
            canvas.configure(scrollregion=canvas.bbox('all'))
        self.mapping_frame.bind('<Configure>', _on_configure)
        
        header_row3 = ttk.Frame(self.mapping_frame)
        header_row3.pack(fill=tk.X, pady=2)
        ttk.Label(header_row3, text="模板字段", width=18, anchor='center', font=('Microsoft YaHei', 8, 'bold')).pack(side=tk.LEFT, padx=3)
        ttk.Label(header_row3, text="数据源字段（请选择）", width=28, anchor='center', font=('Microsoft YaHei', 8, 'bold')).pack(side=tk.LEFT, padx=3)
        
        self.mapping_widgets = {}
        choices = ['不填充'] + self.source_fields
        
        for field in self.template_fields:
            row_frame = ttk.Frame(self.mapping_frame)
            row_frame.pack(fill=tk.X, pady=1)
            
            ttk.Label(row_frame, text=field, width=18, anchor='w').pack(side=tk.LEFT, padx=3)
            
            combo = ttk.Combobox(row_frame, values=choices, width=28, state='readonly')
            combo.pack(side=tk.LEFT, padx=3)
            
            auto_match = self._auto_match(field)
            if auto_match and auto_match in self.source_fields:
                combo.set(auto_match)
            else:
                combo.set('不填充')
            
            self.mapping_widgets[field] = combo
        
        right_frame = ttk.Frame(paned, width=250)
        paned.add(right_frame, weight=0)
        
        settings_frame = ttk.LabelFrame(right_frame, text="导出设置", padding=10)
        settings_frame.pack(fill=tk.BOTH, expand=True)
        
        date_row3 = ttk.Frame(settings_frame)
        date_row3.pack(fill=tk.X, pady=4)
        ttk.Label(date_row3, text="审计基准日:").pack(anchor=tk.W)
        self.base_date_var = tk.StringVar(value=self.base_date)
        ttk.Entry(date_row3, textvariable=self.base_date_var, width=14).pack(anchor=tk.W, pady=(2, 0))
        ttk.Label(date_row3, text="（自动替换日期）", foreground='gray', font=('Microsoft YaHei', 8)).pack(anchor=tk.W)
        
        file_row3 = ttk.Frame(settings_frame)
        file_row3.pack(fill=tk.X, pady=4)
        ttk.Label(file_row3, text="输出文件名:").pack(anchor=tk.W)
        default_filename = f"{Path(self.template_path).stem}_已填充_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        self.output_filename_var = tk.StringVar(value=default_filename)
        ttk.Entry(file_row3, textvariable=self.output_filename_var, width=18).pack(anchor=tk.W, pady=(2, 0))
        
        range_row3 = ttk.Frame(settings_frame)
        range_row3.pack(fill=tk.X, pady=4)
        ttk.Label(range_row3, text="填充范围:").pack(anchor=tk.W)
        self.range_var = tk.StringVar(value="全部数据")
        range_combo = ttk.Combobox(range_row3, textvariable=self.range_var, values=['全部数据', '有差异的数据', '自定义行数'], width=16, state='readonly')
        range_combo.pack(anchor=tk.W, pady=(2, 0))
        
        self.custom_count_frame = ttk.Frame(range_row3)
        ttk.Label(self.custom_count_frame, text="行数:").pack(side=tk.LEFT)
        self.custom_count_var = tk.StringVar(value="50")
        ttk.Entry(self.custom_count_frame, textvariable=self.custom_count_var, width=6).pack(side=tk.LEFT, padx=2)
        self.custom_count_frame.pack(anchor=tk.W, pady=(2, 0))
        self.custom_count_frame.pack_forget()
        
        range_combo.bind('<<ComboboxSelected>>', self._on_range_changed)
        
        self.include_total_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(settings_frame, text="包含合计行", variable=self.include_total_var).pack(anchor=tk.W, pady=2)
        
        self.insert_rows_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(settings_frame, text="数据超出时自动插入行", variable=self.insert_rows_var).pack(anchor=tk.W, pady=2)
        
        btn_frame = ttk.Frame(settings_frame)
        btn_frame.pack(fill=tk.X, pady=10)
        ttk.Button(btn_frame, text="📊 导出", command=self._do_export).pack(side=tk.LEFT, padx=3)
        ttk.Button(btn_frame, text="取消", command=self._on_cancel).pack(side=tk.LEFT, padx=3)
        
        self.status_label = ttk.Label(settings_frame, text="就绪", foreground='gray')
        self.status_label.pack(anchor=tk.W, pady=2)
    
    def _on_range_changed(self, event):
        if self.range_var.get() == '自定义行数':
            self.custom_count_frame.pack(anchor=tk.W, pady=(2, 0))
        else:
            self.custom_count_frame.pack_forget()
    
    def _on_cancel(self):
        self.result = None
        self.window.destroy()
    
    def _auto_match(self, template_field):
        rules = {
            '固定资产名称': ['资产名称', 'asset_name'],
            '取得时间': ['入账日期', 'purchase_date'],
            '使用年限': ['折旧年限', 'useful_life'],
            '固定资产原值': ['资产原值', 'original_cost'],
            '原值': ['资产原值', 'original_cost'],
            '残值率': ['净残值率(%)', 'residual_rate'],
            '累计折旧期初余额': ['账面累计折旧', 'accumulated_dep'],
            '累计折旧': ['账面累计折旧', 'accumulated_dep'],
            '减值准备期初余额': ['减值准备'],
            '本期应提折旧': ['测算本年折旧', 'theory_annual_dep'],
            '本期已提折旧': ['账面本年折旧', 'book_annual_dep'],
            '计提差异': ['本年折旧差异', 'diff_annual'],
            '审计差异': ['累计折旧差异', 'diff_accum'],
            '折旧审定数': ['测算累计折旧', 'theory_accum_dep'],
        }
        for key, targets in rules.items():
            if key in template_field or template_field in key:
                for t in targets:
                    if t in self.source_fields:
                        return t
        return None
    
    def _do_export(self):
        try:
            self.status_label.config(text="正在导出...", foreground='blue')
            self.window.update()
            
            try:
                header_row = int(self.header_entry.get().strip())
            except ValueError:
                messagebox.showerror("错误", "表头行必须为数字")
                return
            
            try:
                data_start_row = int(self.data_start_entry.get().strip())
            except ValueError:
                messagebox.showerror("错误", "数据起始行必须为数字")
                return
            
            total_row = None
            total_str = self.total_entry.get().strip()
            if total_str:
                try:
                    total_row = int(total_str)
                except ValueError:
                    messagebox.showerror("错误", "合计行必须为数字")
                    return
            
            audit_row = None
            audit_str = self.audit_entry.get().strip()
            if audit_str:
                try:
                    audit_row = int(audit_str)
                except ValueError:
                    messagebox.showerror("错误", "审计说明行必须为数字")
                    return
            
            mapping = {}
            for field, combo in self.mapping_widgets.items():
                selected = combo.get()
                if selected and selected != '不填充':
                    mapping[field] = selected
            
            if not mapping:
                messagebox.showwarning("提示", "没有选择任何字段映射")
                return
            
            range_type = self.range_var.get()
            custom_count = None
            if range_type == '自定义行数':
                try:
                    custom_count = int(self.custom_count_var.get().strip() or 50)
                except ValueError:
                    custom_count = 50
            
            data = self.merged_data.copy()
            if data.empty:
                messagebox.showwarning("提示", "没有可用的数据")
                return
            
            if range_type == '有差异的数据':
                cond = []
                for col in ['累计折旧差异', '本年折旧差异']:
                    if col in data.columns:
                        cond.append(abs(data[col]) > 0.01)
                if cond:
                    mask = cond[0]
                    for c in cond[1:]:
                        mask = mask | c
                    data = data[mask]
                else:
                    data = data.head(0)
            elif range_type == '自定义行数' and custom_count:
                data = data.head(custom_count)
            
            if data.empty:
                messagebox.showwarning("提示", "没有符合条件的数据")
                return
            
            wb = load_workbook(self.template_path)
            ws = wb.active
            
            base_date = self.base_date_var.get().strip()
            if base_date:
                date_pattern = r'\d{4}-\d{2}-\d{2}'
                for row in ws.iter_rows(max_row=header_row):
                    for cell in row:
                        if cell.value and isinstance(cell.value, str):
                            if '截止日期' in cell.value or '日期：' in cell.value:
                                if re.search(date_pattern, cell.value):
                                    cell.value = re.sub(date_pattern, base_date, cell.value)
            
            data_start = data_start_row
            
            if total_row is not None:
                end_row = total_row - 1
            elif audit_row is not None:
                end_row = audit_row - 1
            else:
                end_row = ws.max_row
            
            available_rows = end_row - data_start + 1
            data_rows = len(data)
            
            merged_ranges = []
            for merged_cell in ws.merged_cells.ranges:
                merged_ranges.append({
                    'min_row': merged_cell.min_row,
                    'max_row': merged_cell.max_row,
                    'min_col': merged_cell.min_col,
                    'max_col': merged_cell.max_col,
                })
            
            insert_count = 0
            if data_rows > available_rows and total_row is not None and self.insert_rows_var.get():
                insert_count = data_rows - available_rows
                insert_pos = total_row - 1
                if insert_pos > 0:
                    ws.insert_rows(insert_pos, amount=insert_count)
                    total_row += insert_count
                    end_row += insert_count
                    if audit_row is not None:
                        audit_row += insert_count
            
            if insert_count > 0:
                max_col = ws.max_column
                for offset in range(insert_count):
                    target_row = data_start + offset
                    copy_row_style(ws, data_start_row, target_row, max_col)
            
            if insert_count > 0:
                for merged in list(ws.merged_cells.ranges):
                    if merged.min_row > header_row:
                        try:
                            ws.unmerge_cells(str(merged))
                        except Exception:
                            pass
                for mr in merged_ranges:
                    if mr['min_row'] <= header_row:
                        continue
                    new_min = mr['min_row'] + insert_count
                    new_max = mr['max_row'] + insert_count
                    if new_min > new_max or new_max > ws.max_row:
                        continue
                    try:
                        ws.merge_cells(
                            start_row=new_min,
                            start_column=mr['min_col'],
                            end_row=new_max,
                            end_column=mr['max_col']
                        )
                    except Exception:
                        pass
            
            field_cols = {}
            for idx, field in enumerate(self.template_fields, start=1):
                field_cols[field] = idx
            
            for offset, (_, row_data) in enumerate(data.iterrows()):
                row_num = data_start + offset
                if total_row is not None and row_num >= total_row:
                    break
                if audit_row is not None and row_num >= audit_row:
                    break
                
                for template_field, source_field in mapping.items():
                    if template_field not in field_cols:
                        continue
                    col = field_cols[template_field]
                    if source_field in row_data.index:
                        val = row_data[source_field]
                        if isinstance(val, pd.Timestamp):
                            val = val.strftime('%Y-%m-%d')
                        elif isinstance(val, float):
                            val = round(val, 2)
                        ws.cell(row=row_num, column=col, value=val)
            
            if self.include_total_var.get() and total_row is not None:
                for template_field, source_field in mapping.items():
                    if template_field not in field_cols:
                        continue
                    col = field_cols[template_field]
                    total = 0
                    has_value = False
                    for r in range(data_start, total_row):
                        if r >= ws.max_row:
                            break
                        cell_val = ws.cell(row=r, column=col).value
                        if cell_val is not None and isinstance(cell_val, (int, float)):
                            total += cell_val
                            has_value = True
                    if has_value:
                        ws.cell(row=total_row, column=col, value=total)
            
            # ============================================================
            # 修改点：支持“另存为”
            # ============================================================
            output_path = filedialog.asksaveasfilename(
                parent=self.window,
                title="保存模板导出结果",
                defaultextension=".xlsx",
                filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")],
                initialfile=self.output_filename_var.get().strip() or f"{Path(self.template_path).stem}_已填充_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            )
            if not output_path:
                self.status_label.config(text="已取消导出", foreground='orange')
                return
            
            wb.save(output_path)
            
            self.result = output_path
            self.status_label.config(text="✅ 导出完成", foreground='green')
            
            if messagebox.askyesno("导出成功", f"报告已导出至: {output_path}\n是否打开文件？"):
                os.startfile(output_path)
            
            self.window.destroy()
            
        except Exception as e:
            self.status_label.config(text=f"❌ 导出失败: {str(e)[:50]}", foreground='red')
            messagebox.showerror("错误", f"导出失败: {str(e)}")


def export_by_template(parent, df, dep_result, base_date):
    try:
        template_path = filedialog.askopenfilename(
            parent=parent,
            title="选择模板文件",
            filetypes=[("Excel文件", "*.xlsx *.xls"), ("所有文件", "*.*")]
        )
        if not template_path:
            return None
        
        if df is None or df.empty:
            messagebox.showerror("错误", "没有可用的数据，请先加载资产文件")
            return None
        
        settings = ExportSettingsWindow(parent, template_path, df, dep_result, base_date)
        if not settings.initialized:
            settings.window.destroy()
            return None
        
        parent.wait_window(settings.window)
        return settings.result
    except Exception as e:
        messagebox.showerror("导出错误", f"发生异常：{str(e)}")
        return None