# gui.py
# 完整支持 .xls (xlrd) 和 .xlsx (openpyxl)
# 修改1：场景改名（1→程序1，3→程序2，4→程序3），删除场景2
# 修改2：STEP 2 按钮横向一排放在标题下方，状态提示在按钮下方
# 修改3：导出栏从左侧STEP 4移至右侧底部固定一行
# 修改4：集成按模板导出功能（export_template.py）
# 修改5：方案加载后左上角显示方案名称，STEP3复选框同步方案配置
# 修改6：快速导出支持“另存为”（选择路径和文件名）
# 修改7：表头预览只显示前15列，表头行黄色高亮
# 修改8：STEP1增加“启用资产类型自动分类”复选框
# 修改9：添加 import json

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import threading
import pandas as pd
import os
from pathlib import Path
import json  # <-- 添加这一行

from data_loader import load_dataframe, load_data_with_mapping
from calculator import calculate_depreciation
from checker import check_cutoff, check_over_depreciated
from report import generate_reports
from field_matcher import FIELD_DISPLAY_NAMES, REQUIRED_FIELDS, save_mapping, get_all_keywords, load_or_create_mapping

from scheme_manager import (
    get_current_scheme,
    SchemeManager
)
from config import THRESHOLDS

# 导入按模板导出模块
from export_template import export_by_template

# ===== 新增：导入资产分类器 =====
from asset_classifier import show_classifier, load_keyword_library, save_keyword_library


class DepreciationApp:
    def __init__(self, root):
        self.root = root
        root.title("固定资产折旧审计工具 v1.0")
        root.geometry("1400x780")
        root.minsize(1200, 700)

        self.file_path = tk.StringVar()
        self.header_row = tk.StringVar(value="3")
        self.df = None
        self.mapping = {}
        self.excel_columns = []
        self.current_step = 1

        self.mapping_widgets = {}

        self.dep_result = None
        self.cutoff_result = None
        self.over_result = None
        self.analysis_result = None
        self.report_path = None

        self.current_scheme = None

        main_panel = ttk.PanedWindow(root, orient=tk.HORIZONTAL)
        main_panel.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)

        left_frame = ttk.Frame(main_panel, width=420)
        main_panel.add(left_frame, weight=0)
        right_frame = ttk.Frame(main_panel)
        main_panel.add(right_frame, weight=1)

        self._build_left_panel(left_frame)
        self._build_right_panel(right_frame)
        self._build_statusbar(root)

        self._refresh_mapping_frame()

        self._load_startup_scheme()


    # ============================================================
    # 方案管理相关方法
    # ============================================================

    def _load_startup_scheme(self):
        scheme = get_current_scheme()
        if scheme:
            self.current_scheme = scheme
            self._apply_scheme_to_ui(scheme)
            self.status_text.config(text=f"已加载方案: {scheme.get('scheme_name', '未命名')}")
            self.status_file.config(text=f"方案: {scheme.get('scheme_name', '')}")
        else:
            self.status_text.config(text="使用默认配置")
            self.base_date_var.set(THRESHOLDS.get('audit_base_date', '2025-12-31'))

    def _apply_scheme_to_ui(self, scheme):
        """将方案参数应用到界面控件和全局配置"""
        try:
            gp = scheme.get('global_params', {})
            THRESHOLDS['audit_base_date'] = gp.get('audit_base_date', '2025-12-31')
            THRESHOLDS['residual_rate_min'] = gp.get('residual_rate_min', 0)
            THRESHOLDS['residual_rate_max'] = gp.get('residual_rate_max', 5)
            THRESHOLDS['materiality_diff'] = gp.get('materiality_threshold', 500)
            self.base_date_var.set(gp.get('audit_base_date', '2025-12-31'))

            scheme_name = scheme.get('scheme_name', '未命名')
            self.scheme_name_label.config(text=scheme_name)
            self.status_text.config(text=f"方案已加载: {scheme_name}")
            self.status_file.config(text=f"方案: {scheme_name}")

            scenes = scheme.get('scenes', {})
            self.scene1_var.set(scenes.get('scene1', {}).get('enabled', True))
            self.scene3_var.set(scenes.get('scene3', {}).get('enabled', True))
            self.scene4_var.set(scenes.get('scene4', {}).get('enabled', True))

            print(f"✅ 方案已应用: {scheme_name}")
            print(f"   基准日: {THRESHOLDS['audit_base_date']}")
            print(f"   残值率范围: {THRESHOLDS['residual_rate_min']}% ~ {THRESHOLDS['residual_rate_max']}%")
            print(f"   重要性阈值: {THRESHOLDS['materiality_diff']} 元")
            print(f"   程序1: {'启用' if self.scene1_var.get() else '禁用'}")
            print(f"   程序2: {'启用' if self.scene3_var.get() else '禁用'}")
            print(f"   程序3: {'启用' if self.scene4_var.get() else '禁用'}")

        except Exception as e:
            print(f"⚠️ 应用方案到界面时出错: {e}")

    def _apply_scheme_from_manager(self, scheme_data):
        self.current_scheme = scheme_data
        self._apply_scheme_to_ui(scheme_data)
        messagebox.showinfo("方案已应用", f"方案 '{scheme_data.get('scheme_name', '')}' 已应用到主程序！")

    def _open_scheme_manager(self):
        SchemeManager(self.root, on_apply_callback=self._apply_scheme_from_manager)


    # ============================================================
    # 界面构建
    # ============================================================

    def _build_left_panel(self, parent):
        scheme_toolbar = ttk.Frame(parent)
        scheme_toolbar.pack(fill=tk.X, pady=(0, 5))
        ttk.Label(scheme_toolbar, text="📋 方案:", font=('Microsoft YaHei', 9)).pack(side=tk.LEFT)
        self.scheme_name_label = ttk.Label(scheme_toolbar, text="未加载", foreground='gray')
        self.scheme_name_label.pack(side=tk.LEFT, padx=5)
        ttk.Button(scheme_toolbar, text="⚙️ 方案管理", command=self._open_scheme_manager).pack(side=tk.LEFT, padx=5)

        ttk.Label(parent, text="⚙️ 审计配置", font=('Microsoft YaHei', 14, 'bold')).pack(pady=(0, 10))

        # ----- STEP 1 -----
        step1_frame = ttk.LabelFrame(parent, text="STEP 1: 文件与表头设置", padding=10)
        step1_frame.pack(fill=tk.X, pady=5)

        file_row = ttk.Frame(step1_frame)
        file_row.pack(fill=tk.X, pady=2)
        ttk.Label(file_row, text="资产文件:").pack(side=tk.LEFT)
        ttk.Entry(file_row, textvariable=self.file_path).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)
        ttk.Button(file_row, text="浏览", command=self._browse_file).pack(side=tk.LEFT)

        header_row = ttk.Frame(step1_frame)
        header_row.pack(fill=tk.X, pady=5)
        ttk.Label(header_row, text="表头行号(从1开始):").pack(side=tk.LEFT)
        ttk.Entry(header_row, textvariable=self.header_row, width=6).pack(side=tk.LEFT, padx=5)
        ttk.Button(header_row, text="🔍 自动检测", command=self._auto_detect).pack(side=tk.LEFT, padx=2)
        ttk.Button(header_row, text="📋 预览", command=self._preview_data).pack(side=tk.LEFT, padx=2)

        # ===== 新增：分类复选框 + 加载按钮放在一行 =====
        load_row = ttk.Frame(step1_frame)
        load_row.pack(fill=tk.X, pady=5)
        
        self.classify_enabled = tk.BooleanVar(value=False)
        ttk.Checkbutton(load_row, text="启用资产类型自动分类", variable=self.classify_enabled).pack(side=tk.LEFT, padx=2)
        ttk.Button(load_row, text="✅ 加载", command=self._load_data).pack(side=tk.LEFT, padx=2)

        self.step1_status = ttk.Label(step1_frame, text="⏳ 请选择文件", foreground='gray', wraplength=350)
        self.step1_status.pack(anchor=tk.W, pady=2)

        # ----- STEP 2 -----
        step2_frame = ttk.LabelFrame(parent, text="STEP 2: 字段映射确认", padding=10)
        step2_frame.pack(fill=tk.BOTH, expand=True, pady=5)

        map_btn_row = ttk.Frame(step2_frame)
        map_btn_row.pack(fill=tk.X, pady=(0, 5))
        ttk.Button(map_btn_row, text="✅ 确认映射", command=self._confirm_mapping).pack(side=tk.LEFT, padx=2)
        ttk.Button(map_btn_row, text="🔄 重新匹配", command=self._auto_match_mapping).pack(side=tk.LEFT, padx=2)
        ttk.Button(map_btn_row, text="💾 保存映射配置", command=self._save_mapping_config).pack(side=tk.LEFT, padx=2)
        ttk.Button(map_btn_row, text="📂 打开配置文件", command=self._open_mapping_file).pack(side=tk.LEFT, padx=2)

        self.map_canvas = tk.Canvas(step2_frame, height=200)
        scrollbar = ttk.Scrollbar(step2_frame, orient=tk.VERTICAL, command=self.map_canvas.yview)
        self.map_canvas.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.map_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        self.map_inner_frame = ttk.Frame(self.map_canvas)
        self.map_canvas.create_window((0, 0), window=self.map_inner_frame, anchor='nw')
        self.map_inner_frame.bind('<Configure>', lambda e: self.map_canvas.configure(scrollregion=self.map_canvas.bbox('all')))

        header_frame = ttk.Frame(self.map_inner_frame)
        header_frame.pack(fill=tk.X, pady=2)
        ttk.Label(header_frame, text="状态", width=8, anchor='center').pack(side=tk.LEFT, padx=2)
        ttk.Label(header_frame, text="标准字段", width=16, anchor='center').pack(side=tk.LEFT, padx=2)
        ttk.Label(header_frame, text="匹配的Excel列名（请选择）", width=30, anchor='center').pack(side=tk.LEFT, padx=2)

        self.mapping_rows_frame = ttk.Frame(self.map_inner_frame)
        self.mapping_rows_frame.pack(fill=tk.BOTH, expand=True)

        self.step2_status = ttk.Label(step2_frame, text="⏳ 请先加载数据", foreground='gray', wraplength=350)
        self.step2_status.pack(anchor=tk.W, pady=2)

        # ----- STEP 3 -----
        step3_frame = ttk.LabelFrame(parent, text="STEP 3: 执行检查", padding=10)
        step3_frame.pack(fill=tk.X, pady=5)

        check_row1 = ttk.Frame(step3_frame)
        check_row1.pack(anchor=tk.W)
        self.scene1_var = tk.BooleanVar(value=True)
        self.scene3_var = tk.BooleanVar(value=True)
        self.scene4_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(check_row1, text="程序1: 折旧测算差异", variable=self.scene1_var).pack(side=tk.LEFT, padx=5)

        check_row2 = ttk.Frame(step3_frame)
        check_row2.pack(anchor=tk.W)
        ttk.Checkbutton(check_row2, text="程序2: 已提足检查（未提满补提）", variable=self.scene3_var).pack(side=tk.LEFT, padx=5)
        ttk.Checkbutton(check_row2, text="程序3: 异常分析（差异原因诊断）", variable=self.scene4_var).pack(side=tk.LEFT, padx=5)

        base_row = ttk.Frame(step3_frame)
        base_row.pack(fill=tk.X, pady=5)
        ttk.Label(base_row, text="审计基准日:").pack(side=tk.LEFT)
        self.base_date_var = tk.StringVar(value="2025-12-31")
        ttk.Entry(base_row, textvariable=self.base_date_var, width=12).pack(side=tk.LEFT, padx=5)

        exec_row = ttk.Frame(step3_frame)
        exec_row.pack(fill=tk.X, pady=5)
        self.exec_btn = ttk.Button(exec_row, text="▶ 执行选中的检查", command=self._execute_checks)
        self.exec_btn.pack(side=tk.LEFT)

        self.progress_bar = ttk.Progressbar(step3_frame, mode='indeterminate')
        self.progress_bar.pack(fill=tk.X, pady=5)
        self.progress_label = ttk.Label(step3_frame, text="就绪", foreground='gray', wraplength=350)
        self.progress_label.pack(anchor=tk.W)

    def _build_right_panel(self, parent):
        # 右侧整体容器
        right_container = ttk.Frame(parent)
        right_container.pack(fill=tk.BOTH, expand=True)

        top_row = ttk.Frame(right_container)
        top_row.pack(fill=tk.X, pady=(0, 5))
        ttk.Label(top_row, text="📋 数据预览 / 结果展示", font=('Microsoft YaHei', 12, 'bold')).pack(side=tk.LEFT)
        self.info_label = ttk.Label(top_row, text="", foreground='gray')
        self.info_label.pack(side=tk.RIGHT)

        self.notebook = ttk.Notebook(right_container)
        self.notebook.pack(fill=tk.BOTH, expand=True)

        self.preview_tab = ttk.Frame(self.notebook)
        self.notebook.add(self.preview_tab, text="📄 表头预览")
        self._build_preview_tab(self.preview_tab)

        self.dep_tab = ttk.Frame(self.notebook)
        self.notebook.add(self.dep_tab, text="📊 程序1_折旧测算差异")
        self._build_result_tab(self.dep_tab, 'dep')

        self.checker_tab = ttk.Frame(self.notebook)
        self.notebook.add(self.checker_tab, text="📋 程序2_已提足检查")
        self._build_result_tab(self.checker_tab, 'checker')

        self.analysis_tab = ttk.Frame(self.notebook)
        self.notebook.add(self.analysis_tab, text="🔍 程序3_异常分析")
        self._build_result_tab(self.analysis_tab, 'analysis')

        self.summary_tab = ttk.Frame(self.notebook)
        self.notebook.add(self.summary_tab, text="📈 汇总统计")
        self._build_summary_tab(self.summary_tab)

        # ----- 底部固定导出栏 -----
        export_frame = ttk.Frame(right_container)
        export_frame.pack(side=tk.BOTTOM, fill=tk.X, pady=(5, 0))

        export_row = ttk.Frame(export_frame)
        export_row.pack(fill=tk.X)
        ttk.Button(export_row, text="📊 快速导出", command=self._export_excel).pack(side=tk.LEFT, padx=2)
        ttk.Button(export_row, text="📄 按模板导出", command=self._export_by_template).pack(side=tk.LEFT, padx=2)
        ttk.Button(export_row, text="📂 打开输出文件夹", command=self._open_output_folder).pack(side=tk.LEFT, padx=2)
        self.export_status = ttk.Label(export_row, text="未导出", foreground='gray')
        self.export_status.pack(side=tk.LEFT, padx=10)

    def _build_preview_tab(self, parent):
        container = ttk.Frame(parent)
        container.pack(fill=tk.BOTH, expand=True)
        self.preview_tree = ttk.Treeview(container, show='headings')
        self.preview_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar = ttk.Scrollbar(container, orient=tk.VERTICAL, command=self.preview_tree.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.preview_tree.configure(yscrollcommand=scrollbar.set)
        h_scrollbar = ttk.Scrollbar(parent, orient=tk.HORIZONTAL, command=self.preview_tree.xview)
        h_scrollbar.pack(fill=tk.X)
        self.preview_tree.configure(xscrollcommand=h_scrollbar.set)

    def _build_result_tab(self, parent, tab_type):
        container = ttk.Frame(parent)
        container.pack(fill=tk.BOTH, expand=True)
        count_label = ttk.Label(container, text="共 0 条", foreground='gray')
        count_label.pack(anchor=tk.W, pady=2)
        setattr(self, f'{tab_type}_count_label', count_label)

        tree = ttk.Treeview(container, show='headings')
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar = ttk.Scrollbar(container, orient=tk.VERTICAL, command=tree.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        tree.configure(yscrollcommand=scrollbar.set)
        h_scrollbar = ttk.Scrollbar(container, orient=tk.HORIZONTAL, command=tree.xview)
        h_scrollbar.pack(fill=tk.X)
        tree.configure(xscrollcommand=h_scrollbar.set)
        setattr(self, f'{tab_type}_tree', tree)

    def _build_summary_tab(self, parent):
        self.summary_text = tk.Text(parent, font=('Consolas', 11), wrap=tk.WORD)
        self.summary_text.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        self.summary_text.insert('1.0', '等待执行检查...')
        self.summary_text.config(state='disabled')

    def _build_statusbar(self, parent):
        status_frame = ttk.Frame(parent)
        status_frame.pack(fill=tk.X, side=tk.BOTTOM, padx=5, pady=2)
        self.status_icon = ttk.Label(status_frame, text="●", foreground='green')
        self.status_icon.pack(side=tk.LEFT)
        self.status_text = ttk.Label(status_frame, text="就绪", foreground='gray')
        self.status_text.pack(side=tk.LEFT, padx=5)
        ttk.Separator(status_frame, orient=tk.VERTICAL).pack(side=tk.LEFT, padx=10, fill=tk.Y)
        self.status_file = ttk.Label(status_frame, text="未加载文件", foreground='gray')
        self.status_file.pack(side=tk.LEFT, padx=5)
        ttk.Separator(status_frame, orient=tk.VERTICAL).pack(side=tk.LEFT, padx=10, fill=tk.Y)
        self.status_count = ttk.Label(status_frame, text="资产数: 0", foreground='gray')
        self.status_count.pack(side=tk.LEFT, padx=5)


    # ============================================================
    # 功能方法
    # ============================================================

    def _browse_file(self):
        filename = filedialog.askopenfilename(
            filetypes=[("Excel文件", "*.xlsx *.xls"), ("所有文件", "*.*")]
        )
        if filename:
            self.file_path.set(filename)
            self.step1_status.config(text="📁 文件已选择，点击'加载'查看", foreground='blue')
            self._update_status("文件已选择", f"文件: {Path(filename).name}")

    def _auto_detect(self):
        file_path = self.file_path.get().strip()
        if not file_path or not Path(file_path).exists():
            messagebox.showerror("错误", "请先选择有效的Excel文件")
            return
        try:
            ext = Path(file_path).suffix.lower()
            if ext == '.xls':
                import xlrd
                wb = xlrd.open_workbook(file_path)
                sheet = wb.sheet_by_index(0)
                all_keywords = get_all_keywords()
                best_score = 0
                best_row = 1
                for row_idx in range(min(20, sheet.nrows)):
                    row_vals = []
                    for col_idx in range(sheet.ncols):
                        cv = sheet.cell_value(row_idx, col_idx)
                        if cv:
                            row_vals.append(str(cv))
                    score = 0
                    for cell in row_vals:
                        for kw in all_keywords:
                            if kw in cell:
                                score += 1
                                break
                    if score > best_score:
                        best_score = score
                        best_row = row_idx + 1
                if best_score >= 3:
                    self.header_row.set(str(best_row))
                    self.step1_status.config(text=f"✅ 自动检测到表头行: 第{best_row}行", foreground='green')
                else:
                    messagebox.showwarning("提示", "无法自动检测表头行，请手动指定")
            else:
                try:
                    from openpyxl import load_workbook
                    wb = load_workbook(file_path, data_only=True)
                    ws = wb.active
                    all_keywords = get_all_keywords()
                    best_score = 0
                    best_row = 1
                    for row_idx in range(1, 21):
                        row_vals = []
                        for col_idx in range(1, ws.max_column + 1):
                            cv = ws.cell(row=row_idx, column=col_idx).value
                            if cv is not None:
                                row_vals.append(str(cv))
                        score = 0
                        for cell in row_vals:
                            for kw in all_keywords:
                                if kw in cell:
                                    score += 1
                                    break
                        if score > best_score:
                            best_score = score
                            best_row = row_idx
                    if best_score >= 3:
                        self.header_row.set(str(best_row))
                        self.step1_status.config(text=f"✅ 自动检测到表头行: 第{best_row}行", foreground='green')
                    else:
                        messagebox.showwarning("提示", "无法自动检测表头行，请手动指定")
                except Exception as e:
                    messagebox.showerror("错误", f"自动检测失败: {e}")
        except Exception as e:
            messagebox.showerror("错误", f"自动检测失败: {e}")

    def _preview_data(self):
        file_path = self.file_path.get().strip()
        if not file_path or not Path(file_path).exists():
            messagebox.showerror("错误", "请先选择有效的Excel文件")
            return
        try:
            header = int(self.header_row.get()) if self.header_row.get().isdigit() else 3
            df_full = load_dataframe(file_path, header)
            if df_full is None or df_full.empty:
                messagebox.showerror("错误", "无法读取文件数据")
                return

            df_preview = df_full.head(15).copy()
            df_preview.insert(0, '行号', range(1, len(df_preview) + 1))
            self._update_preview_tree(df_preview)

            self.step1_status.config(
                text=f"✅ 预览完成，共{len(df_preview)}行，表头行: 第{header}行（高亮显示）",
                foreground='green'
            )
        except ImportError as e:
            messagebox.showerror("缺少依赖", str(e))
            self.step1_status.config(text="❌ 缺少依赖", foreground='red')
        except Exception as e:
            messagebox.showerror("预览异常", f"发生异常:\n{str(e)}")
            self.step1_status.config(text=f"❌ 异常: {str(e)[:50]}", foreground='red')

    def _update_preview_tree(self, df_preview):
        tree = self.preview_tree
        for item in tree.get_children():
            tree.delete(item)

        max_cols = 15
        cols = list(df_preview.columns)
        if len(cols) > max_cols:
            display_cols = cols[:max_cols]
        else:
            display_cols = cols

        tree['columns'] = display_cols

        for col in display_cols:
            tree.heading(col, text=str(col))
            tree.column(col, width=100, minwidth=80)

        header_row_num = int(self.header_row.get()) if self.header_row.get().isdigit() else 3

        for idx, row in df_preview.iterrows():
            values = [row[col] for col in display_cols]
            item = tree.insert('', 'end', values=values)

            row_num = row.get('行号', 0)
            if row_num == header_row_num:
                tree.tag_configure('header', background='#FFD700')
                tree.item(item, tags=('header',))

    def _load_data(self):
        """加载数据（修改：支持资产类型自动分类）"""
        file_path = self.file_path.get().strip()
        if not file_path or not Path(file_path).exists():
            messagebox.showerror("错误", "请先选择有效的Excel文件")
            return
        header = int(self.header_row.get()) if self.header_row.get().isdigit() else 3

        try:
            df_full = load_dataframe(file_path, header)
            if df_full is None or df_full.empty:
                messagebox.showerror("错误", "无法读取文件数据")
                self._refresh_mapping_frame()
                return

            self.excel_columns = [col for col in df_full.columns if col != '' and col is not None]

            if not self.excel_columns:
                messagebox.showerror("读取列名失败", "未能从Excel中读取到任何列名，请检查表头行是否正确。")
                self.step1_status.config(text="❌ 列名为空", foreground='red')
                self._refresh_mapping_frame()
                return

            print(f"✅ 读取到 {len(self.excel_columns)} 个列名：{self.excel_columns[:10]}...")

            self.mapping = load_or_create_mapping(self.excel_columns)
            if self.mapping is None:
                messagebox.showinfo("提示", "已生成 mapping_config.json，请检查并重新加载")
                self.step1_status.config(text="⚠️ 已生成映射配置文件，请检查后重新加载", foreground='orange')
                self._refresh_mapping_frame()
                return

            # ===== 先加载数据 =====
            self.df = load_data_with_mapping(file_path, header, self.mapping)

            # ===== 新增：如果启用了分类，弹出分类确认窗口 =====
            if self.classify_enabled.get():
                # 加载API配置（如果有）
                api_config = self._load_api_config()
                
                # 弹出分类窗口
                result_df = show_classifier(self.root, self.df, api_config)
                
                if result_df is None:
                    # 用户取消了分类，取消本次加载
                    self.step1_status.config(text="⏳ 已取消加载", foreground='orange')
                    return
                
                # 使用分类后的DataFrame
                self.df = result_df
                
                # 更新列名列表，包含新增的分类列
                for new_col in ['资产类型', '税法最低年限', '分类来源']:
                    if new_col in self.df.columns and new_col not in self.excel_columns:
                        self.excel_columns.append(new_col)
                
                self.step1_status.config(
                    text=f"✅ 加载成功！共{len(self.df)}条资产数据（已自动分类）", 
                    foreground='green'
                )
            else:
                self.step1_status.config(
                    text=f"✅ 加载成功！共{len(self.df)}条资产数据", 
                    foreground='green'
                )

            self._update_status("数据已加载", f"文件: {Path(file_path).name}")
            self.status_count.config(text=f"资产数: {len(self.df)}")
            self._update_info_label()
            self.current_step = 2

        except ImportError as e:
            messagebox.showerror("缺少依赖", str(e))
            self.step1_status.config(text="❌ 缺少依赖", foreground='red')
        except ValueError as e:
            messagebox.showerror("加载失败", f"错误详情:\n{str(e)}")
            self.step1_status.config(text=f"❌ {str(e)[:50]}...", foreground='red')
        except Exception as e:
            messagebox.showerror("加载失败", f"错误详情:\n{str(e)}")
            self.step1_status.config(text=f"❌ {str(e)[:50]}...", foreground='red')
        finally:
            self._refresh_mapping_frame()
            # 刷新映射框架时，新列会出现在下拉框中

    def _load_api_config(self):
        """加载API配置"""
        config_file = Path("./api_config.json")
        if config_file.exists():
            try:
                with open(config_file, 'r', encoding='utf-8') as f:
                    return json.load(f)
            except Exception:
                pass
        return {'enabled': False}

    def _save_api_config(self, config):
        """保存API配置"""
        config_file = Path("./api_config.json")
        with open(config_file, 'w', encoding='utf-8') as f:
            json.dump(config, f, ensure_ascii=False, indent=4)

    # ============================================================
    # 以下方法保持不变
    # ============================================================

    def _refresh_mapping_frame(self):
        for widget in self.mapping_rows_frame.winfo_children():
            widget.destroy()
        self.mapping_widgets.clear()

        choices = ['请选择...'] + self.excel_columns if self.excel_columns else ['请选择...']

        for field in FIELD_DISPLAY_NAMES.keys():
            display = FIELD_DISPLAY_NAMES[field]
            is_required = ' *' if field in REQUIRED_FIELDS else ''
            row_frame = ttk.Frame(self.mapping_rows_frame)
            row_frame.pack(fill=tk.X, pady=1)

            status_label = ttk.Label(row_frame, text="●", foreground='gray', width=8, anchor='center')
            status_label.pack(side=tk.LEFT, padx=2)
            setattr(status_label, '_field', field)

            ttk.Label(row_frame, text=f"{display}{is_required}", width=16, anchor='w').pack(side=tk.LEFT, padx=2)

            combo = ttk.Combobox(row_frame, values=choices, width=30, state='readonly')
            combo.pack(side=tk.LEFT, padx=2, fill=tk.X, expand=True)
            combo.set('请选择...')
            combo.bind('<<ComboboxSelected>>', lambda e, f=field: self._on_mapping_select(f))
            self.mapping_widgets[field] = combo

        if self.mapping:
            for field, col_name in self.mapping.items():
                if field in self.mapping_widgets and col_name in self.excel_columns:
                    self.mapping_widgets[field].set(col_name)

        self._update_mapping_status()

    def _on_mapping_select(self, field):
        combo = self.mapping_widgets.get(field)
        if combo:
            val = combo.get()
            if val != '请选择...':
                self.mapping[field] = val
            else:
                if field in self.mapping:
                    del self.mapping[field]
        self._update_mapping_status()

    def _auto_match_mapping(self):
        if not self.excel_columns:
            messagebox.showwarning("提示", "请先加载数据")
            return
        from field_matcher import smart_match_columns
        auto = smart_match_columns(self.excel_columns)
        for field, col in auto.items():
            if field in self.mapping_widgets and col in self.excel_columns:
                self.mapping_widgets[field].set(col)
                self.mapping[field] = col
        self._update_mapping_status()
        self.step2_status.config(text="✅ 自动匹配完成，请检查并确认", foreground='green')

    def _update_mapping_status(self):
        missing = []
        for field in REQUIRED_FIELDS:
            combo = self.mapping_widgets.get(field)
            if combo:
                val = combo.get()
                if val != '请选择...' and val in self.excel_columns:
                    for child in self.mapping_rows_frame.winfo_children():
                        for subchild in child.winfo_children():
                            if isinstance(subchild, ttk.Label) and hasattr(subchild, '_field') and subchild._field == field:
                                subchild.config(foreground='green')
                                break
                else:
                    for child in self.mapping_rows_frame.winfo_children():
                        for subchild in child.winfo_children():
                            if isinstance(subchild, ttk.Label) and hasattr(subchild, '_field') and subchild._field == field:
                                subchild.config(foreground='red')
                                break
                    if val == '请选择...' or val not in self.excel_columns:
                        missing.append(FIELD_DISPLAY_NAMES.get(field, field))

        if missing:
            self.step2_status.config(
                text=f"⚠️ 缺少必填字段: {', '.join(missing)}，请从下拉框选择",
                foreground='red'
            )
            self.step1_status.config(
                text=f"⚠️ 缺少必填字段: {', '.join(missing)}",
                foreground='red'
            )
        else:
            self.step2_status.config(text="✅ 所有必填字段已匹配", foreground='green')
            self.step1_status.config(
                text="✅ 所有必填字段已匹配，请点击'确认映射'",
                foreground='green'
            )

    def _confirm_mapping(self):
        missing = []
        for field in REQUIRED_FIELDS:
            combo = self.mapping_widgets.get(field)
            if combo:
                val = combo.get()
                if val == '请选择...' or val not in self.excel_columns:
                    missing.append(FIELD_DISPLAY_NAMES.get(field, field))
            else:
                missing.append(FIELD_DISPLAY_NAMES.get(field, field))

        if missing:
            messagebox.showerror("错误", f"以下必填字段未选择有效的列名: {', '.join(missing)}")
            return

        for field, combo in self.mapping_widgets.items():
            val = combo.get()
            if val != '请选择...' and val in self.excel_columns:
                self.mapping[field] = val
            else:
                if field in self.mapping:
                    del self.mapping[field]

        save_mapping(self.mapping)
        self.step2_status.config(text="✅ 映射已确认并保存", foreground='green')
        self.step1_status.config(text="✅ 映射已确认并保存，请进入 STEP 3 执行检查", foreground='green')
        messagebox.showinfo("成功", "字段映射已确认！请进入STEP 3执行检查")

    def _save_mapping_config(self):
        for field, combo in self.mapping_widgets.items():
            val = combo.get()
            if val != '请选择...' and val in self.excel_columns:
                self.mapping[field] = val
        save_mapping(self.mapping)
        messagebox.showinfo("成功", "映射配置已保存到 mapping_config.json")

    def _open_mapping_file(self):
        mapping_path = Path('mapping_config.json')
        if mapping_path.exists():
            os.startfile(mapping_path)
        else:
            messagebox.showinfo("提示", "配置文件不存在")

    def _update_info_label(self):
        if self.df is not None:
            self.info_label.config(text=f"共 {len(self.df)} 条资产 | 表头行: 第{self.header_row.get()}行")

    def _update_status(self, text, detail=""):
        self.status_text.config(text=text)
        if detail:
            self.status_file.config(text=detail)

    def _execute_checks(self):
        if self.df is None or self.df.empty:
            messagebox.showerror("错误", "请先加载资产数据")
            return

        missing = []
        for field in REQUIRED_FIELDS:
            if field not in self.mapping or not self.mapping[field]:
                missing.append(FIELD_DISPLAY_NAMES.get(field, field))
        if missing:
            messagebox.showerror("错误", f"以下必填字段未映射: {', '.join(missing)}")
            return

        THRESHOLDS['audit_base_date'] = self.base_date_var.get().strip()

        self.exec_btn.config(state='disabled')
        self.progress_bar.start()
        self.progress_label.config(text="执行中...", foreground='blue')
        self._update_status("执行中...")
        thread = threading.Thread(target=self._run_checks)
        thread.daemon = True
        thread.start()

    def _run_checks(self):
        try:
            self._update_progress("执行程序1: 折旧测算...")

            if self.scene1_var.get():
                self.dep_result = calculate_depreciation(self.df)
            else:
                self.dep_result = None

            self._update_progress("执行程序2: 已提足检查...")
            if self.scene3_var.get():
                self.over_result = check_over_depreciated(self.df)
            else:
                self.over_result = pd.DataFrame()

            self._update_progress("执行程序3: 异常分析...")
            if self.scene4_var.get() and self.dep_result is not None and not self.dep_result.empty:
                from analyzer import analyze_differences
                self.analysis_result = analyze_differences(self.dep_result)
            else:
                self.analysis_result = pd.DataFrame()

            self.cutoff_result = pd.DataFrame()

            self.root.after(0, self._display_results)
        except Exception as e:
            self.root.after(0, lambda: self._show_error(str(e)))

    def _update_progress(self, msg):
        self.root.after(0, lambda: self.progress_label.config(text=msg))

    def _display_results(self):
        self._show_result_in_tab('dep', self.dep_result)
        self._show_result_in_tab('checker', self.over_result)
        self._show_result_in_tab('analysis', self.analysis_result)
        self._update_summary()

        self.progress_bar.stop()
        self.progress_label.config(text="✅ 执行完成", foreground='green')
        self.exec_btn.config(state='normal')
        self._update_status("执行完成")
        self.notebook.select(0)
        messagebox.showinfo("完成", "检查执行完成！请查看结果标签页")

    def _show_result_in_tab(self, tab_type, df):
        tree = getattr(self, f'{tab_type}_tree')
        count_label = getattr(self, f'{tab_type}_count_label')
        for item in tree.get_children():
            tree.delete(item)
        if df is None or df.empty:
            count_label.config(text="共 0 条")
            return
        cols = list(df.columns)
        tree['columns'] = cols
        for col in cols:
            tree.heading(col, text=col)
            tree.column(col, width=140)
        for idx, row in df.head(10000).iterrows():
            values = [row[col] for col in cols]
            tree.insert('', 'end', values=values)
        count_label.config(text=f"共 {len(df)} 条")

    def _update_summary(self):
        summary_text = self.summary_text
        summary_text.config(state='normal')
        summary_text.delete('1.0', 'end')
        lines = []
        lines.append("=" * 50)
        lines.append("📊 审计检查汇总统计")
        lines.append("=" * 50)
        lines.append("")
        if self.df is not None:
            lines.append(f"📌 资产总数: {len(self.df)}")
        else:
            lines.append("📌 资产总数: 0")
        lines.append("")
        lines.append("【程序1: 折旧测算差异】")
        if self.dep_result is not None and not self.dep_result.empty:
            total_diff = self.dep_result['累计折旧差异'].sum()
            diff_count = len(self.dep_result[abs(self.dep_result['累计折旧差异']) > 0.5])
            material_count = len(self.dep_result[abs(self.dep_result['累计折旧差异']) > THRESHOLDS.get('materiality_diff', 500)])
            if '净残值校验状态' in self.dep_result.columns:
                check_fail_count = len(self.dep_result[self.dep_result['净残值校验状态'] == '已提足校验异常'])
            else:
                check_fail_count = 0
            lines.append(f"  差异总额: {total_diff:.2f} 元")
            lines.append(f"  有差异资产数: {diff_count} 条")
            lines.append(f"  重大差异(>{THRESHOLDS.get('materiality_diff', 500)}元): {material_count} 条")
            lines.append(f"  已提足校验异常: {check_fail_count} 条")
        else:
            lines.append("  未执行或数据为空")
        lines.append("")
        lines.append("【程序2: 已提足检查（未提满补提）】")
        if self.over_result is not None and not self.over_result.empty:
            lines.append(f"  已提足未提满资产数: {len(self.over_result)} 条")
        else:
            lines.append("  未执行或数据为空")
        lines.append("")
        lines.append("【程序3: 异常分析】")
        if self.analysis_result is not None and not self.analysis_result.empty:
            lines.append(f"  有诊断结论资产数: {len(self.analysis_result)} 条")
            if '诊断结论' in self.analysis_result.columns:
                multi = len(self.analysis_result[self.analysis_result['诊断结论'].str.contains('多提', na=False)])
                short = len(self.analysis_result[self.analysis_result['诊断结论'].str.contains('少提', na=False)])
                lines.append(f"  多提折旧资产: {multi} 条")
                lines.append(f"  少提折旧资产: {short} 条")
        else:
            lines.append("  未执行或数据为空")
        lines.append("")
        lines.append("=" * 50)
        lines.append("✅ 检查完成，请查看右侧标签页")
        summary_text.insert('1.0', '\n'.join(lines))
        summary_text.config(state='disabled')

    def _show_error(self, msg):
        self.progress_bar.stop()
        self.exec_btn.config(state='normal')
        self.progress_label.config(text=f"❌ 错误: {msg[:50]}", foreground='red')
        self._update_status("错误", msg[:30])
        messagebox.showerror("执行错误", msg)

    def _export_excel(self):
        """快速导出（支持选择保存路径和文件名）"""
        has_data = (self.dep_result is not None and not self.dep_result.empty) or \
                   (self.over_result is not None and not self.over_result.empty) or \
                   (self.analysis_result is not None and not self.analysis_result.empty)
        if not has_data:
            messagebox.showwarning("提示", "没有可导出的结果，请先执行检查")
            return

        output_path = filedialog.asksaveasfilename(
            parent=self.root,
            title="保存报告",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")],
            initialfile="折旧测算差异报告.xlsx"
        )
        if not output_path:
            return

        try:
            generate_reports(self.dep_result, self.over_result, self.analysis_result, output_path)
            self.report_path = output_path
            self.export_status.config(text=f"✅ 已导出: {Path(output_path).name}", foreground='green')
            self._update_status("导出成功", Path(output_path).name)
            messagebox.showinfo("成功", f"报告已导出至: {output_path}")
        except Exception as e:
            messagebox.showerror("错误", f"导出失败: {e}")

    def _export_by_template(self):
        from export_template import export_by_template
        if self.df is None or self.df.empty:
            messagebox.showerror("错误", "没有可用的数据，请先加载资产文件")
            return
        result = export_by_template(
            parent=self.root,
            df=self.df,
            dep_result=self.dep_result,
            base_date=THRESHOLDS.get('audit_base_date', '2025-12-31')
        )
        if result:
            self.export_status.config(text=f"✅ 已导出: {Path(result).name}", foreground='green')
            self._update_status("导出成功", Path(result).name)

    def _open_output_folder(self):
        os.startfile(Path('.'))


def run_gui():
    root = tk.Tk()
    app = DepreciationApp(root)
    root.mainloop()


if __name__ == '__main__':
    run_gui()